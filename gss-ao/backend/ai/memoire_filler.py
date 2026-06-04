"""Remplissage du template imposé Univ Rouen (Mode A) via python-docx.

On part du fichier ORIGINAL `5-Mémoire Technique - lots 1-2-3.docx` (copié dans
`templates/`) et on remplit, en place, les zones de réponse (pointillés) repérées
par l'intitulé de leur question — sans reproduire le document (mise en forme
d'origine préservée). Les sections générées par l'IA proviennent du frontend.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from backend.ai.sections import SECTIONS_BY_ID

TEMPLATE_PATH = Path(__file__).parent / "templates" / "memoire-univ-rouen-template.docx"

_DOTS_RE = re.compile(r"[.…\s_]")


def _is_dotted(text: str) -> bool:
    """True si le paragraphe est essentiellement une zone de pointillés à remplir."""
    t = text.strip()
    if not t:
        return False
    return len(_DOTS_RE.sub("", t)) < 3


def _set_text(p: Paragraph, text: str) -> None:
    """Remplace le texte d'un paragraphe en conservant le style du 1er run."""
    runs = p.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def _fill_inline(p: Paragraph, value: str) -> None:
    """Zone 'intitulé : ....' sur une seule ligne -> conserve l'intitulé, remplace
    les pointillés par la valeur."""
    head = p.text.split(":")[0].rstrip()
    _set_text(p, f"{head} : {value}")


def _norm(s: str) -> str:
    """Normalise pour le matching : minuscules + apostrophes droites."""
    return s.lower().replace("’", "'").replace("‘", "'")


def _find_label_index(paras: list[Paragraph], label: str) -> int | None:
    low = _norm(label)
    for i, p in enumerate(paras):
        if low in _norm(p.text):
            return i
    return None


def _fill_zone(paras: list[Paragraph], label: str, value: str, *, inline: bool) -> bool:
    """Remplit la zone de réponse associée à `label`. Retourne True si rempli."""
    idx = _find_label_index(paras, label)
    if idx is None:
        return False
    if inline:
        _fill_inline(paras[idx], value)
        return True
    # zone 'next' : 1er paragraphe pointillés dans les 5 suivants
    for j in range(idx + 1, min(idx + 6, len(paras))):
        if _is_dotted(paras[j].text):
            _set_text(paras[j], value)
            # nettoyer les pointillés résiduels consécutifs de la même zone
            for k in range(j + 1, min(j + 6, len(paras))):
                if _is_dotted(paras[k].text):
                    _set_text(paras[k], "")
                else:
                    break
            return True
    return False


def fill_template(
    sections: dict[str, str],
    *,
    identite: dict | None = None,
    signataire: str | None = None,
    date_signature: str | None = None,
) -> tuple[bytes, dict]:
    """Remplit le template et retourne (docx_bytes, rapport).

    Args:
        sections: {section_id -> texte généré}.
        identite: {denomination, num_cnaps, date_autorisation} (données fixes).
        signataire / date_signature: bloc signature.

    Returns:
        (bytes du .docx, rapport {filled: [...], missing: [...]}).
    """
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template introuvable : {TEMPLATE_PATH}")

    doc = Document(str(TEMPLATE_PATH))
    paras = doc.paragraphs
    filled: list[str] = []
    missing: list[str] = []

    # --- Identité (zones inline en haut du document) ---
    identite = identite or {}
    id_zones = [
        ("Dénomination du candidat", identite.get("denomination")),
        ("N° CNAPS d", identite.get("num_cnaps")),
        ("Date d'obtention de l", identite.get("date_autorisation")),
    ]
    for label, value in id_zones:
        if value and _fill_zone(paras, label, value, inline=True):
            filled.append(label)

    # --- Sections rédigées par l'IA ---
    for section_id, text in sections.items():
        spec = SECTIONS_BY_ID.get(section_id)
        if spec is None or not (text and text.strip()):
            continue
        ok = _fill_zone(paras, spec.label, text.strip(), inline=spec.inline)
        (filled if ok else missing).append(section_id)

    # --- Signature ---
    if signataire or date_signature:
        sig_idx = _find_label_index(paras, "Date et signature du candidat")
        if sig_idx is not None:
            val = " ".join(x for x in [date_signature, "—", signataire] if x)
            _fill_inline(paras[sig_idx], val)
            filled.append("signature")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), {"filled": filled, "missing": missing}
