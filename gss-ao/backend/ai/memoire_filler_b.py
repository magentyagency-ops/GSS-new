"""MODE B — construction du mémoire technique GSS (réponse libre).

AO RNE.docx (Clarence) est un mémoire FINI (≈2770 paragraphes), pas un formulaire
à trous : on s'appuie sur sa STRUCTURE (I. Présentation / II. Moyens humains /
III. Moyens opérationnels / IV. Moyens organisationnels) pour produire, via
python-docx, un mémoire GSS propre et éditable rempli par l'IA, avec en annexe
les slides SLIDE REP AO sélectionnées.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from backend.ai.sections_b import CHAPITRES_B, SECTIONS_B

_INDIGO = RGBColor(0x43, 0x38, 0xCA)
_SLATE = RGBColor(0x33, 0x41, 0x55)


def build_memoire_b(
    sections: dict[str, str],
    *,
    selected_slides: list[dict] | None = None,
    projet: str = "Mémoire technique",
    acheteur: str = "",
    candidat: str = "GSS — Sécurité privée",
    signataire: str | None = None,
    date_signature: str | None = None,
) -> tuple[bytes, dict]:
    """Retourne (docx_bytes, rapport)."""
    doc = Document()
    doc.styles["Normal"].font.name = "Helvetica"
    doc.styles["Normal"].font.size = Pt(10)

    # --- Page de titre ---
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("MÉMOIRE TECHNIQUE")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = _INDIGO
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(projet + (f"\n{acheteur}" if acheteur else ""))
    sr.font.size = Pt(13)
    cand = doc.add_paragraph()
    cand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cand.add_run(f"Candidat : {candidat}")
    cr.bold = True
    doc.add_page_break()

    filled: list[str] = []
    missing: list[str] = []

    # --- Chapitres I → IV ---
    for chap in ["I", "II", "III", "IV"]:
        h = doc.add_heading(level=1)
        hr = h.add_run(f"{chap}. {CHAPITRES_B[chap].upper()}")
        hr.font.color.rgb = _INDIGO
        for spec in [s for s in SECTIONS_B if s.chapter == chap]:
            doc.add_heading(spec.title, level=2)
            text = (sections.get(spec.id) or "").strip()
            if text:
                for para in text.split("\n\n"):
                    doc.add_paragraph(para.strip())
                filled.append(spec.id)
            else:
                p = doc.add_paragraph()
                pr = p.add_run("[Section à générer]")
                pr.italic = True
                pr.font.color.rgb = _SLATE
                missing.append(spec.id)

    # --- Annexe : slides sélectionnées ---
    slides = selected_slides or []
    if slides:
        doc.add_page_break()
        doc.add_heading("Annexe — Slides GSS mobilisées", level=1)
        for chap in ["I", "II", "III", "IV"]:
            chap_slides = [s for s in slides if s.get("chapter") == chap]
            if not chap_slides:
                continue
            doc.add_heading(f"{chap}. {CHAPITRES_B[chap]}", level=2)
            for s in chap_slides:
                doc.add_paragraph(
                    f"{s.get('category', '')} — {s.get('file_name', '')}", style="List Bullet"
                )

    # --- Signature ---
    if signataire or date_signature:
        doc.add_paragraph()
        sig = doc.add_paragraph()
        sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        parts = [f"Fait le {date_signature}" if date_signature else "", signataire or ""]
        sig.add_run(" — ".join(x for x in parts if x))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), {"filled": filled, "missing": missing, "slides": len(slides)}
